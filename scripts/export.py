#!/usr/bin/env python3
"""
scripts/export.py
Reads data/orders.json and generates:
  - exports/orders.xlsx (with Orders sheet & Summary sheet)
  - exports/orders.csv (UTF-8 with BOM)
  - exports/orders.md (Markdown report)

Can be run standalone or imported by app.py.
"""
import os
import json
import csv
import pathlib
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
DATA_PATH = ROOT / "data" / "orders.json"
EXPORTS_DIR = ROOT / "exports"

def load_orders():
    if not DATA_PATH.exists():
        return []
    with open(DATA_PATH, "r", encoding="utf-8") as f:
        try:
            orders = json.load(f)
            if isinstance(orders, list):
                return orders
            return []
        except Exception as e:
            print(f"Error reading {DATA_PATH}: {e}")
            return []

def safe_number(val, default=0):
    if val is None or val == "":
        return default
    try:
        return float(val)
    except (ValueError, TypeError):
        return default

def export_all():
    """Generates XLSX, CSV, and MD exports from current data/orders.json"""
    EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    orders = load_orders()
    
    # Sort orders by expected delivery date (earliest first, blanks at end)
    sorted_orders = sorted(
        orders,
        key=lambda o: (o.get("expectedDeliveryDate") or o.get("expected") or "9999-99-99")
    )

    xlsx_path = EXPORTS_DIR / "orders.xlsx"
    csv_path = EXPORTS_DIR / "orders.csv"
    md_path = EXPORTS_DIR / "orders.md"

    export_xlsx(sorted_orders, xlsx_path)
    export_csv(sorted_orders, csv_path)
    export_md(sorted_orders, md_path)

    print(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] Exports generated successfully:")
    print(f"  - {xlsx_path}")
    print(f"  - {csv_path}")
    print(f"  - {md_path}")
    return {"xlsx": str(xlsx_path), "csv": str(csv_path), "md": str(md_path)}

def export_xlsx(orders, filepath):
    wb = Workbook()
    
    # --- Sheet 1: Orders ---
    ws_orders = wb.active
    ws_orders.title = "Orders"

    headers = [
        "Order ID", "Product Name", "Platform", "Order Number",
        "Price (INR)", "Quantity", "Total (INR)", "Payment Method",
        "Order Date", "Expected Delivery", "Actual Delivery",
        "Status", "Received", "Notes"
    ]
    ws_orders.append(headers)

    # Styles
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid") # Deep Teal
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color="E2E8F0"),
        right=Side(style='thin', color="E2E8F0"),
        top=Side(style='thin', color="E2E8F0"),
        bottom=Side(style='thin', color="E2E8F0")
    )

    for col_idx, cell in enumerate(ws_orders[1], 1):
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center

    for row_idx, o in enumerate(orders, start=2):
        price = safe_number(o.get("price") if o.get("price") is not None else o.get("amount"))
        qty = int(o.get("quantity") or 1)
        total = price * qty if (o.get("price") is not None or o.get("amount") is not None) else 0
        received_str = "Yes" if o.get("received") else "No"
        
        row_values = [
            o.get("id", f"ORD-{row_idx-1}"),
            o.get("productName") or o.get("item", ""),
            o.get("platform", "Meesho"),
            o.get("orderNo", ""),
            price if (o.get("price") is not None or o.get("amount") is not None) else None,
            qty,
            f"=E{row_idx}*F{row_idx}",
            o.get("paymentMethod") or o.get("payment", "Cash on Delivery"),
            o.get("orderedDate") or o.get("ordered", ""),
            o.get("expectedDeliveryDate") or o.get("expected", ""),
            o.get("actualDeliveryDate") or "",
            o.get("status", "Order Placed"),
            received_str,
            o.get("notes", "")
        ]
        ws_orders.append(row_values)

        for col_idx, cell in enumerate(ws_orders[row_idx], 1):
            cell.font = Font(name="Segoe UI", size=10)
            cell.border = thin_border
            if col_idx in (5, 7): # Price & Total
                cell.number_format = '"₹"#,##0.00'
                cell.alignment = align_right
            elif col_idx in (1, 6, 8, 9, 10, 11, 12, 13):
                cell.alignment = align_center
            elif col_idx == 4: # Order No
                cell.number_format = "@"
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    total_rows = len(orders) + 1
    col_widths = [14, 36, 16, 24, 15, 12, 16, 18, 14, 18, 16, 18, 12, 35]
    for i, w in enumerate(col_widths, 1):
        ws_orders.column_dimensions[get_column_letter(i)].width = w

    ws_orders.freeze_panes = "A2"
    ws_orders.auto_filter.ref = f"A1:N{max(total_rows, 2)}"

    # --- Sheet 2: Summary ---
    ws_summary = wb.create_sheet("Summary")
    ws_summary.append(["Order Tracker - KPI Summary", ""])
    ws_summary.append(["Metric", "Value"])

    # Header style for summary
    ws_summary.merge_cells("A1:B1")
    ws_summary["A1"].font = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
    ws_summary["A1"].fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Navy
    ws_summary["A1"].alignment = align_center

    for col_idx, col_name in enumerate(["A2", "B2"], 1):
        ws_summary[col_name].font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        ws_summary[col_name].fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
        ws_summary[col_name].alignment = align_center

    summary_rows = [
        ("Total Orders", f"=COUNTA(Orders!A2:A{total_rows})"),
        ("Delivered / Received Orders", f'=COUNTIF(Orders!M2:M{total_rows}, "Yes")'),
        ("Pending Deliveries", f'=COUNTIF(Orders!M2:M{total_rows}, "No")'),
        ("Orders On the Way", f'=COUNTIF(Orders!L2:L{total_rows}, "On the Way")'),
        ("Orders Out for Delivery", f'=COUNTIF(Orders!L2:L{total_rows}, "Out for Delivery")'),
        ("Cancelled / Returned", f'=COUNTIF(Orders!L2:L{total_rows}, "Cancelled") + COUNTIF(Orders!L2:L{total_rows}, "Returned")'),
        ("Total Order Value", f'=SUM(Orders!G2:G{total_rows})'),
        ("Pending Order Value", f'=SUMIFS(Orders!G2:G{total_rows}, Orders!M2:M{total_rows}, "No")'),
        ("Received Order Value", f'=SUMIFS(Orders!G2:G{total_rows}, Orders!M2:M{total_rows}, "Yes")'),
        ("Generated On", datetime.now().strftime("%d %b %Y, %I:%M %p"))
    ]

    for label, formula in summary_rows:
        ws_summary.append([label, formula])

    for row_idx in range(3, 3 + len(summary_rows)):
        c1 = ws_summary.cell(row=row_idx, column=1)
        c2 = ws_summary.cell(row=row_idx, column=2)
        c1.font = Font(name="Segoe UI", size=10, bold=True)
        c2.font = Font(name="Segoe UI", size=10)
        c1.border = thin_border
        c2.border = thin_border
        if row_idx in (9, 10, 11): # Values with currency
            c2.number_format = '"₹"#,##0.00'
            c2.alignment = align_right
        else:
            c2.alignment = align_center

    ws_summary.column_dimensions["A"].width = 30
    ws_summary.column_dimensions["B"].width = 25

    note_cell = ws_summary.cell(row=14, column=1, value="Note: The formulas dynamically calculate values when opened in Microsoft Excel.")
    note_cell.font = Font(name="Segoe UI", size=9, italic=True, color="64748B")

    wb.save(filepath)

def export_csv(orders, filepath):
    headers = [
        "Order ID", "Product Name", "Platform", "Order Number",
        "Price (INR)", "Quantity", "Total (INR)", "Payment Method",
        "Order Date", "Expected Delivery", "Actual Delivery",
        "Status", "Received", "Notes"
    ]
    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for o in orders:
            price = safe_number(o.get("price") if o.get("price") is not None else o.get("amount"))
            qty = int(o.get("quantity") or 1)
            total = price * qty if (o.get("price") is not None or o.get("amount") is not None) else ""
            writer.writerow([
                o.get("id", ""),
                o.get("productName") or o.get("item", ""),
                o.get("platform", "Meesho"),
                o.get("orderNo", ""),
                price if (o.get("price") is not None or o.get("amount") is not None) else "",
                qty,
                total,
                o.get("paymentMethod") or o.get("payment", "Cash on Delivery"),
                o.get("orderedDate") or o.get("ordered", ""),
                o.get("expectedDeliveryDate") or o.get("expected", ""),
                o.get("actualDeliveryDate") or "",
                o.get("status", "Order Placed"),
                "Yes" if o.get("received") else "No",
                o.get("notes", "")
            ])

def export_md(orders, filepath):
    now_str = datetime.now().strftime("%d %B %Y, %I:%M %p")
    total_orders = len(orders)
    received_count = sum(1 for o in orders if o.get("received"))
    pending_count = total_orders - received_count
    on_the_way_count = sum(1 for o in orders if o.get("status") in ("On the Way", "Out for Delivery"))
    cancelled_count = sum(1 for o in orders if o.get("status") in ("Cancelled", "Returned"))
    
    total_value = sum(
        safe_number(o.get("price") if o.get("price") is not None else o.get("amount")) * int(o.get("quantity") or 1)
        for o in orders if (o.get("price") is not None or o.get("amount") is not None)
    )
    pending_value = sum(
        safe_number(o.get("price") if o.get("price") is not None else o.get("amount")) * int(o.get("quantity") or 1)
        for o in orders if not o.get("received") and (o.get("price") is not None or o.get("amount") is not None)
    )

    md = f"""# 📦 Order Tracker Report

> **Generated on:** {now_str}  
> **Source:** `data/orders.json`

---

## 📊 Summary KPI Overview

| Metric | Value |
|:-------|------:|
| **Total Orders** | `{total_orders}` |
| **Delivered / Received** | `{received_count}` |
| **Pending Deliveries** | `{pending_count}` |
| **Orders On The Way** | `{on_the_way_count}` |
| **Cancelled / Returned** | `{cancelled_count}` |
| **Total Order Value** | **₹{total_value:,.2f}** |
| **Pending Amount** | **₹{pending_value:,.2f}** |

---

## 📋 Complete Orders List

| # | Product Name | Platform | Order Number | Price | Qty | Total | Expected By | Status | Received | Notes |
|:-:|:---|:---:|:---:|---:|:-:|---:|:---:|:---:|:---:|:---|
"""
    for idx, o in enumerate(orders, 1):
        price_val = o.get("price") if o.get("price") is not None else o.get("amount")
        if price_val is not None and price_val != "":
            price_str = f"₹{safe_number(price_val):,.0f}"
            qty_val = int(o.get("quantity") or 1)
            total_str = f"₹{safe_number(price_val) * qty_val:,.0f}"
        else:
            price_str = "*Not entered*"
            qty_val = int(o.get("quantity") or 1)
            total_str = "—"

        rec_icon = "✅ Yes" if o.get("received") else "⬜ No"
        order_no = f"`{o.get('orderNo')}`" if o.get('orderNo') else "—"
        prod_name = o.get("productName") or o.get("item", "—")
        platform = o.get("platform", "Meesho")
        exp_date = o.get("expectedDeliveryDate") or o.get("expected") or "—"
        status = o.get("status", "Order Placed")
        notes = o.get("notes") or "—"

        md += f"| {idx} | **{prod_name}** | {platform} | {order_no} | {price_str} | {qty_val} | {total_str} | {exp_date} | `{status}` | {rec_icon} | {notes} |\n"

    md += """
---

*Report automatically generated by Order Tracker System.*
"""
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(md)

if __name__ == "__main__":
    export_all()
